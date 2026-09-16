"""
build_tracker.py
-----------------
One-time (or occasional) full rebuild of Sheet1 in the CA Equity Markets
Tracker - replacing the old Bloomberg BDH pull with a clean, calendar-
complete history sourced from this pipeline. 'Indices Return', 'Tracker',
'Manual' and 'ngen' are left completely untouched: same formulas, same
formatting, same cell positions. They reference Sheet1 by name and by
fixed column position, not by anything Bloomberg-specific, so nothing
else needs to change for them to keep working.

Why rebuild Sheet1 at all, instead of just appending forever:
  - The old row-15 Bloomberg array formulas are dead (cached as #NAME?
    already) and confirmed capable of wiping the real historical values
    below them if this file is ever recalculated without a live
    Bloomberg connection. A clean rebuild has none of those formulas.
  - ~10 sector columns in the old file were tracking price return under
    a "TRI" label. This pipeline's data replaces them with genuine total
    return (see indices_config.csv notes).
  - Columns get readable labels instead of Bloomberg tickers.

What's preserved, per column, to keep maximum history depth:
  - Before this pipeline's own earliest cached date for an index: the
    OLD file's values are kept as-is (real history the pipeline doesn't
    have yet, e.g. Nifty 50 back to 2001).
  - From the pipeline's earliest cached date onward: pipeline data wins,
    forward-filled across weekends/holidays exactly like the old
    convention (so 'Indices Return's exact-match date lookups keep
    working for every possible Analysis Date).
  - The 4 tickers this pipeline doesn't fetch (Nifty 50 TRI, Nifty Next
    50 TRI, Nifty 100 TRI, Nifty 500 TRI) keep their full old-file
    history and simply stop updating at the old file's last row.

After this one-time rebuild, go back to using export_to_tracker.py for
the weekly NAV update - it already appends in this same clean format.

Usage:
    python build_tracker.py --config indices_config.csv \\
        --old-tracker "CA_Equity_Markets_Tracker.xlsx" \\
        --output "CA_Equity_Markets_Tracker_v2.xlsx"
"""
import argparse
import bisect
from datetime import datetime, timedelta

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string

import indices_lib as lib

# Same column map as export_to_tracker.py - the fixed Sheet1 layout that
# 'Indices Return' locks onto by position.
SHEET1_COLUMN_ORDER = [
    ("C", None, "Nifty 50 TRI"),
    ("D", "Nifty 50", "Nifty 50"),
    ("E", "Sensex", "Sensex"),
    ("F", None, "Nifty Next 50 TRI"),
    ("G", "Nifty Next 50", "Nifty Next 50"),
    ("H", None, "Nifty 100 TRI"),
    ("I", "Nifty 100", "Nifty 100"),
    ("J", "Nifty Midcap 150", "Nifty Midcap 150"),
    ("K", "Nifty Large Midcap 250", "Nifty Large Midcap 250"),
    ("L", "Nifty Smallcap 250", "Nifty Smallcap 250"),
    ("M", "Nifty50 Equal Weight", "Nifty50 Equal Weight"),
    ("N", "Nifty100 Equal Weight", "Nifty100 Equal Weight"),
    ("O", None, "Nifty 500 TRI"),
    ("P", "Nifty 200 TRI", "Nifty 200 TRI"),
    ("Q", "NSE 500", "NSE 500"),
    ("R", "Nifty200 Value 30 TRI", "Nifty200 Value 30 TRI"),
    ("S", "Nifty 200 Momentum 30 Index TRI", "Nifty 200 Momentum 30 Index TRI"),
    ("T", "Nifty 200 Quality 30 Index TRI", "Nifty 200 Quality 30 Index TRI"),
    ("U", "Nifty 100 Low Volatility 30", "Nifty 100 Low Volatility 30"),
    ("V", "Nifty 200 Momentum 30 Index", "Nifty 200 Momentum 30 Index"),
    ("W", "Nifty 100 Quality 30", "Nifty 100 Quality 30"),
    ("X", "Nifty Alpha 50", "Nifty Alpha 50"),
    ("Y", "Nifty Alpha Low Volatility 30", "Nifty Alpha Low Volatility 30"),
    ("Z", "Nifty India Defence TRI", "Nifty India Defence TRI"),
    ("AA", "Nifty PSU Bank TRI", "Nifty PSU Bank TRI"),
    ("AB", "Nifty PSE TRI", "Nifty PSE TRI"),
    ("AC", "Nifty Realty TRI", "Nifty Realty TRI"),
    ("AD", "Nifty Auto TRI", "Nifty Auto TRI"),
    ("AE", "Nifty India Manufacturing TRI", "Nifty India Manufacturing TRI"),
    ("AF", "Nifty Infrastructure TRI", "Nifty Infrastructure TRI"),
    ("AG", "Nifty Healthcare TRI", "Nifty Healthcare TRI"),
    ("AH", "Nifty Metal TRI", "Nifty Metal TRI"),
    ("AI", "BSE Consumer Discretionary TRI", "BSE Consumer Discretionary TRI"),
    ("AJ", "Nifty Financial Services TRI", "Nifty Financial Services TRI"),
    ("AK", "Nifty Commodities TRI", "Nifty Commodities TRI"),
    ("AL", "Nifty Bank TRI", "Nifty Bank TRI"),
    ("AM", "BSE Utilities TRI", "BSE Utilities TRI"),
    ("AN", "Nifty Private Bank TRI", "Nifty Private Bank TRI"),
    ("AO", "NIFTY100 ESG TRI", "NIFTY100 ESG TRI"),
    ("AP", "Nifty FMCG TRI", "Nifty FMCG TRI"),
    ("AQ", "Nifty Oil & Gas TRI", "Nifty Oil & Gas TRI"),
    ("AR", "Nifty Energy TRI", "Nifty Energy TRI"),
    ("AS", "Nifty IT TRI", "Nifty IT TRI"),
    ("AT", "S&P 500(US)", "S&P 500(US)"),
    ("AU", "Nasdaq(US)", "Nasdaq(US)"),
    ("AV", "Nikkei(Japan)", "Nikkei(Japan)"),
    ("AW", "Dax Index(Germany)", "Dax Index(Germany)"),
    ("AX", "CAC 40 Index(France)", "CAC 40 Index(France)"),
    ("AY", "FTSE(UK)", "FTSE(UK)"),
    ("AZ", "Hang Seng(Hong Kong)", "Hang Seng(Hong Kong)"),
    ("BA", "USD/INR", "USD/INR"),
    ("BB", "Gold(USD)", "Gold(USD)"),
    ("BC", "Silver(USD)", "Silver(USD)"),
    ("BD", "Shanghai(China)", "Shanghai(China)"),
    ("BE", "BSE 500", "BSE 500"),
    ("BF", "Nifty MidSmall 400", "Nifty MidSmall 400"),
]

DATA_START_ROW = 15  # matches 'Indices Return's hardcoded $B$15 table start


def load_old_sheet1_values(old_tracker_path):
    """Per column letter -> {date: value}, straight from the old file's
    Sheet1, skipping error placeholders like '#N/A N/A' / '#NAME?'."""
    wb = openpyxl.load_workbook(old_tracker_path, data_only=True)
    ws = wb["Sheet1"]
    by_col = {letter: {} for letter, _, _ in SHEET1_COLUMN_ORDER}
    for r in range(16, ws.max_row + 1):  # row 15 is the old formula row - skip it
        d = ws.cell(row=r, column=2).value
        if not isinstance(d, datetime):
            continue
        d = d.date()
        for letter, _, _ in SHEET1_COLUMN_ORDER:
            v = ws.cell(row=r, column=column_index_from_string(letter)).value
            if isinstance(v, (int, float)):
                by_col[letter][d] = float(v)
    return by_col


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="indices_config.csv")
    parser.add_argument("--history-dir", default=lib.HISTORY_DIR)
    parser.add_argument("--old-tracker", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    config = lib.load_config(args.config)
    label_to_config = {e["label"]: e for e in config}

    print("Reading old Sheet1 (pre-pipeline history to preserve) ...")
    old_by_col = load_old_sheet1_values(args.old_tracker)

    print("Loading this pipeline's cached history ...")
    pipeline_by_label = {}
    for _, label, _ in SHEET1_COLUMN_ORDER:
        if label is None:
            continue
        if label not in label_to_config:
            raise SystemExit(f"'{label}' is in SHEET1_COLUMN_ORDER but missing from {args.config}.")
        pipeline_by_label[label] = lib.load_history(label, args.history_dir)

    # cutover date per covered column = pipeline's own earliest cached date
    cutover = {label: s[0][0].date() for label, s in pipeline_by_label.items() if s}

    all_dates = set()
    for d in list(old_by_col.values()):
        all_dates.update(d.keys())
    for s in pipeline_by_label.values():
        all_dates.update(dt.date() for dt, _ in s)
    if not all_dates:
        raise SystemExit("No data found anywhere - check --old-tracker and --history-dir.")
    start_date, end_date = min(all_dates), max(all_dates)
    print(f"Master date range: {start_date} -> {end_date} "
          f"({(end_date - start_date).days + 1} calendar days)")

    # For the 4 uncovered columns, precompute a sorted (date, value) list so
    # we can forward-fill them too, same convention as everything else.
    # Reason: VLOOKUP treats a truly blank cell as 0, so once the Analysis
    # Date moves past this column's last real update, some period columns
    # in 'Indices Return' would show a nonsensical -100% instead of a flat
    # 0% - forward-filling the last known value avoids that.
    old_sorted_by_col = {
        letter: sorted(d_v.items()) for letter, d_v in old_by_col.items()
    }
    old_dates_by_col = {
        letter: [dd for dd, _ in pairs] for letter, pairs in old_sorted_by_col.items()
    }

    def old_value_on_or_before(col_letter, d):
        dates = old_dates_by_col[col_letter]
        i = bisect.bisect_right(dates, d) - 1
        if i < 0:
            return None
        return old_sorted_by_col[col_letter][i][1]

    def value_for(col_letter, label, d):
        if label is None:
            return old_value_on_or_before(col_letter, d)
        co = cutover.get(label)
        if co is not None and d >= co:
            found = lib.value_on_or_before(pipeline_by_label[label], datetime(d.year, d.month, d.day))
            return found[1] if found else None
        # before pipeline coverage (or pipeline has nothing for this label at all)
        v = old_by_col[col_letter].get(d)
        if v is not None:
            return v
        if co is None:
            return None
        # no old value on this exact date but pipeline exists later - nothing to show
        return None

    wb = openpyxl.load_workbook(args.old_tracker)  # formulas mode - keeps every other sheet intact
    if "Sheet1" in wb.sheetnames:
        del wb["Sheet1"]
    ws = wb.create_sheet("Sheet1", 0)

    ws["B1"] = "CA Equity Markets Tracker - Index History"
    ws["B1"].font = Font(bold=True, size=12)
    ws["B3"] = "Source"
    ws["C3"] = "Automated pipeline (jugaad-data / yfinance / bse / akshare) + preserved pre-2014 Bloomberg history"
    ws["B4"] = "Rebuilt on"
    ws["C4"] = datetime.today()
    ws["C4"].number_format = "dd-mm-yyyy"
    ws["B12"] = "Index Name"
    ws["B14"] = "Date"
    for letter, label, display_name in SHEET1_COLUMN_ORDER:
        c = column_index_from_string(letter)
        ws.cell(row=12, column=c, value=display_name if label else f"{display_name} (not covered)").font = Font(bold=True)
        ws.cell(row=14, column=c, value="Value")

    n_days = (end_date - start_date).days + 1
    r = DATA_START_ROW
    d = start_date
    written = 0
    while d <= end_date:
        date_cell = ws.cell(row=r, column=2, value=datetime(d.year, d.month, d.day))
        date_cell.number_format = "dd-mm-yyyy"
        for letter, label, _ in SHEET1_COLUMN_ORDER:
            v = value_for(letter, label, d)
            if v is not None:
                ws.cell(row=r, column=column_index_from_string(letter), value=round(v, 4))
        r += 1
        written += 1
        d += timedelta(days=1)

    ws.column_dimensions["B"].width = 16
    for letter, _, _ in SHEET1_COLUMN_ORDER:
        ws.column_dimensions[letter].width = 20
    ws.freeze_panes = "C15"

    wb.move_sheet("Sheet1", offset=-wb.sheetnames.index("Sheet1"))
    wb.save(args.output)
    print(f"\nWrote {written} calendar-day rows (rows {DATA_START_ROW}-{r - 1}) -> {args.output}")
    print("'Indices Return', 'Tracker', 'Manual' and 'ngen' were not touched.")


if __name__ == "__main__":
    main()