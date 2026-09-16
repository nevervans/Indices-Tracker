"""
full_report.py
---------------
One workbook, two sheets:
  - "History": raw NAV/index values, one column per index, one row per date
    (full available history - 10+ years once refresh_history.py has backfilled)
  - "Returns": the standard 1W/MTD/.../10Y return table, as of any date

Both come straight from the history/ cache built by refresh_history.py -
no network calls, instant.

Usage:
    python full_report.py --config indices_config.csv --output full_report.xlsx
    python full_report.py --config indices_config.csv --output full_report.xlsx --asof 2026-06-15
    python full_report.py --config indices_config.csv --output full_report.xlsx --start 2015-01-01
"""
import argparse
from datetime import datetime

import openpyxl
from openpyxl.styles import Font

import indices_lib as lib


def build_history_sheet(wb, config, history_dir, start=None, end=None):
    labels = [e["label"] for e in config if e["source"] != "manual"]

    series_by_label = {}
    all_dates = set()
    for label in labels:
        series = lib.load_history(label, history_dir)
        if start:
            series = [(d, v) for d, v in series if d >= start]
        if end:
            series = [(d, v) for d, v in series if d <= end]
        series_by_label[label] = dict(series)
        all_dates.update(d for d, _ in series)

    dates_sorted = sorted(all_dates)

    ws = wb.active
    ws.title = "History"

    ws.cell(row=1, column=1, value="Date").font = Font(bold=True)
    for c, label in enumerate(labels, start=2):
        ws.cell(row=1, column=c, value=label).font = Font(bold=True)

    for r, d in enumerate(dates_sorted, start=2):
        date_cell = ws.cell(row=r, column=1, value=d)
        date_cell.number_format = "yyyy-mm-dd"
        for c, label in enumerate(labels, start=2):
            v = series_by_label[label].get(d)
            if v is not None:
                ws.cell(row=r, column=c, value=round(v, 4))

    ws.column_dimensions["A"].width = 12
    for c in range(2, len(labels) + 2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 20
    ws.freeze_panes = "B2"

    return len(dates_sorted), len(labels)


def build_returns_sheet(wb, config, history_dir, asof):
    labels = [e["label"] for e in config if e["source"] != "manual"]
    columns = ["Index", "As Of Date", "Value"] + [h for h, _ in lib.PERIOD_COLUMNS]

    ws = wb.create_sheet("Returns")
    for c, h in enumerate(columns, start=1):
        ws.cell(row=1, column=c, value=h).font = Font(bold=True)

    for r, label in enumerate(labels, start=2):
        series = lib.load_history(label, history_dir)
        ws.cell(row=r, column=1, value=label)
        if not series:
            for c in range(2, len(columns) + 1):
                ws.cell(row=r, column=c, value="no data")
            continue

        latest = lib.value_on_or_before(series, asof)
        returns = lib.compute_standard_returns(series, asof)

        ws.cell(row=r, column=2, value=latest[0].strftime("%Y-%m-%d") if latest else "-")
        val_cell = ws.cell(row=r, column=3)
        if latest:
            val_cell.value = round(latest[1], 2)
            val_cell.number_format = "#,##0.00"
        else:
            val_cell.value = "-"

        for c, (header, _) in enumerate(lib.PERIOD_COLUMNS, start=4):
            ret = returns.get(header)
            cell = ws.cell(row=r, column=c)
            if ret is None:
                cell.value = "-"
            else:
                cell.value = round(ret, 4)
                cell.number_format = "0.00%"

    ws.column_dimensions["A"].width = 32
    for c in range(2, len(columns) + 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 14
    ws.freeze_panes = "B2"

    return len(labels)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="indices_config.csv")
    parser.add_argument("--history-dir", default=lib.HISTORY_DIR)
    parser.add_argument("--output", required=True)
    parser.add_argument("--asof", help="YYYY-MM-DD for the Returns sheet - defaults to today")
    parser.add_argument("--start", help="YYYY-MM-DD - earliest date to include on the History sheet")
    parser.add_argument("--end", help="YYYY-MM-DD - latest date to include on the History sheet")
    args = parser.parse_args()

    config = lib.load_config(args.config)
    asof = datetime.strptime(args.asof, "%Y-%m-%d") if args.asof else datetime.today()
    start = datetime.strptime(args.start, "%Y-%m-%d") if args.start else None
    end = datetime.strptime(args.end, "%Y-%m-%d") if args.end else None

    wb = openpyxl.Workbook()
    n_dates, n_labels = build_history_sheet(wb, config, args.history_dir, start=start, end=end)
    n_returns = build_returns_sheet(wb, config, args.history_dir, asof)

    wb.save(args.output)
    print(f"History sheet: {n_dates} dates x {n_labels} indices")
    print(f"Returns sheet: {n_returns} indices, as of {asof.strftime('%Y-%m-%d')}")
    print(f"Saved -> {args.output}")


if __name__ == "__main__":
    main()