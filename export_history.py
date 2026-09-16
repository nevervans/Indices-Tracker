"""
export_history.py
------------------
Exports the actual historical NAV/index levels (not returns) to Excel,
straight from the history/ cache built by refresh_history.py. One sheet,
one column per index, one row per date - same shape as the original
workbook's Sheet1, minus the Bloomberg formulas.

Usage:
    # Every index, full available history
    python export_history.py --config indices_config.csv --output history_all.xlsx

    # One index only
    python export_history.py --config indices_config.csv --output nifty50_history.xlsx --index "Nifty 50"

    # A specific window
    python export_history.py --config indices_config.csv --output history_2024_2025.xlsx --start 2024-01-01 --end 2025-12-31
"""
import argparse
from datetime import datetime

import openpyxl
from openpyxl.styles import Font

import indices_lib as lib


def build_history_workbook(config, history_dir, start=None, end=None, only_label=None):
    labels = [only_label] if only_label else [e["label"] for e in config if e["source"] != "manual"]

    series_by_label = {}
    all_dates = set()
    for label in labels:
        series = lib.load_history(label, history_dir)
        if start:
            series = [(d, v) for d, v in series if d >= start]
        if end:
            series = [(d, v) for d, v in series if d <= end]
        series_by_label[label] = dict(series)
        all_dates.update(d for d, _ in series)

    dates_sorted = sorted(all_dates)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "History"

    ws.cell(row=1, column=1, value="Date").font = Font(bold=True)
    for c, label in enumerate(labels, start=2):
        ws.cell(row=1, column=c, value=label).font = Font(bold=True)

    for r, d in enumerate(dates_sorted, start=2):
        date_cell = ws.cell(row=r, column=1, value=d)
        date_cell.number_format = "yyyy-mm-dd"
        for c, label in enumerate(labels, start=2):
            v = series_by_label[label].get(d)
            if v is not None:
                ws.cell(row=r, column=c, value=round(v, 4))

    ws.column_dimensions["A"].width = 12
    for c in range(2, len(labels) + 2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = 20
    ws.freeze_panes = "B2"

    return wb, len(dates_sorted), len(labels)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", default="indices_config.csv")
    parser.add_argument("--history-dir", default=lib.HISTORY_DIR)
    parser.add_argument("--output", required=True)
    parser.add_argument("--index", help="Exact label from indices_config.csv - omit for every index")
    parser.add_argument("--start", help="YYYY-MM-DD - omit for earliest available")
    parser.add_argument("--end", help="YYYY-MM-DD - omit for latest available")
    args = parser.parse_args()

    config = lib.load_config(args.config)
    if args.index and args.index not in [e["label"] for e in config]:
        parser.error(f"'{args.index}' is not a label in {args.config}.")

    start = datetime.strptime(args.start, "%Y-%m-%d") if args.start else None
    end = datetime.strptime(args.end, "%Y-%m-%d") if args.end else None

    wb, n_dates, n_labels = build_history_workbook(
        config, args.history_dir, start=start, end=end, only_label=args.index
    )
    wb.save(args.output)
    print(f"Saved {n_dates} dates x {n_labels} indices -> {args.output}")


if __name__ == "__main__":
    main()