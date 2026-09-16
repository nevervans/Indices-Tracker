"""
update_indices.py
------------------
The standard weekly run: fetch fresh data for every index and write a
Tracker sheet with the usual return periods, as of today (or a given date).
This is what the GitHub Action runs on schedule.

If you want historical NAVs saved for reuse, or returns for an arbitrary
date/custom window, use refresh_history.py + query_returns.py instead -
they're faster on repeat runs since they don't re-fetch everything each time.

Usage:
    python update_indices.py --config indices_config.csv --output tracker_auto.xlsx
    python update_indices.py --config indices_config.csv --output tracker_auto.xlsx --asof 2026-09-04
"""
import argparse
import time
from datetime import date, datetime, timedelta

import openpyxl
from openpyxl.styles import Font

import indices_lib as lib


def build_workbook(config, asof: datetime, delay_sec: float, years_back: int = 12):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tracker"

    ws["C2"] = "CA Equity Markets Tracker"
    ws["C2"].font = Font(bold=True, size=14)
    ws["B5"] = "Analysis Date"
    ws["C5"] = asof

    headers = ["Index Name"] + [h for h, _ in lib.PERIOD_COLUMNS]
    for col, text in enumerate(headers, start=2):
        cell = ws.cell(row=9, column=col, value=text)
        cell.font = Font(bold=True)

    row = 10
    current_group = None
    from_date = date.today() - timedelta(days=365 * years_back)
    to_date = date.today()

    for entry in config:
        label = entry["label"]
        group = entry["group"]
        if group != current_group:
            ws.cell(row=row, column=2, value=group).font = Font(bold=True, italic=True)
            current_group = group
            row += 1

        if entry["source"] == "manual":
            ws.cell(row=row, column=2, value=label)
            ws.cell(row=row, column=3, value="manual - see README")
            print(f"[manual] {label}: not available from a free source, skipping.")
            row += 1
            continue

        print(f"Fetching {label} ({entry['source']}: {entry['symbol']}) ...")
        try:
            series = lib.fetch_series(entry, from_date, to_date)
        except Exception as exc:
            print(f"  [error] {label}: {exc}")
            ws.cell(row=row, column=2, value=label)
            ws.cell(row=row, column=3, value=f"error: {exc}")
            row += 1
            time.sleep(delay_sec)
            continue

        returns = lib.compute_standard_returns(series, asof)
        ws.cell(row=row, column=2, value=label)
        for col, (header, _) in enumerate(lib.PERIOD_COLUMNS, start=3):
            val = returns.get(header)
            cell = ws.cell(row=row, column=col)
            if val is None:
                cell.value = "-"
            else:
                cell.value = round(val, 4)
                cell.number_format = "0.00%"
        row += 1
        time.sleep(delay_sec)

    for col_idx in range(2, 3 + len(lib.PERIOD_COLUMNS)):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 14
    ws.column_dimensions["B"].width = 32

    return wb


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--asof", help="YYYY-MM-DD, defaults to today")
    parser.add_argument("--delay", type=float, default=0.3,
                         help="Seconds to wait between requests (be polite to free sources)")
    args = parser.parse_args()

    asof = datetime.strptime(args.asof, "%Y-%m-%d") if args.asof else datetime.today()
    config = lib.load_config(args.config)

    wb = build_workbook(config, asof, args.delay)
    wb.save(args.output)
    print(f"\nDone -> {args.output}")


if __name__ == "__main__":
    main()